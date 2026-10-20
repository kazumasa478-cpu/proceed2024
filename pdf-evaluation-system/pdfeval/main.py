"""ストレスチェック調査票 読取・集計システム

  python -m pdfeval forms   名簿から本人用QRコード入り調査票PDFを作成
  python -m pdfeval run     inbox のスキャンPDFを読取 → 本人別フォルダに保存 → 集計Excel・結果通知書PDFを作成
  python -m pdfeval todo    要確認の人（読取に自信がない回答が残っている人）と読取不可のページを表示
"""
from __future__ import annotations

import argparse
import json
import logging
import shutil
from pathlib import Path

from . import __version__
from .excel_report import write_summary
from .forms import Person, make_forms, read_roster, valid_id, write_roster
from .omr import IMAGE_SUFFIXES, read_pdf
from .organizer import (load_answer_sheet, person_folder, safe_name, save_page_pdf, set_answer,
                        update_answers)
from .questionnaire import SECTION_BY_KEY, item_label, score
from .report_pdf import combine, write_report

log = logging.getLogger("pdfeval")
BASE_DIR = Path(__file__).resolve().parent.parent


# 以前の版の見本の文言。設定がこのままなら新しい既定値に置き換える
_OLD_PLACEHOLDERS = {
    ("report", "contact"): "相談窓口：〇〇部 健康管理担当（内線 0000）",
    ("report", "interview_request"): "申出先：〇〇部 健康管理担当（内線 0000）",
}


def load_config(path: Path) -> dict:
    """設定を読む。新しい版で増えた項目は既定値で補い、旧版の見本の文言は新しい既定値に更新する。"""
    cfg = json.loads(path.read_text(encoding="utf-8"))
    default_path = BASE_DIR / "config.json"
    if not default_path.exists() or default_path.resolve() == path.resolve():
        return cfg
    defaults = json.loads(default_path.read_text(encoding="utf-8"))
    changed = False
    for section, values in defaults.items():
        if isinstance(values, dict):
            target = cfg.setdefault(section, {})
            for key, value in values.items():
                if key not in target or target[key] == _OLD_PLACEHOLDERS.get((section, key)):
                    target[key] = value
                    changed = True
        elif section not in cfg:
            cfg[section] = values
            changed = True
    if changed:
        path.write_text(json.dumps(cfg, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return cfg


def _roster(base: Path, cfg: dict) -> dict[str, Person]:
    path = base / cfg["roster"]
    if not path.exists():
        raise SystemExit(f"名簿が見つかりません: {path}（列: ID,氏名,所属 のCSV）")
    try:
        return {p.id: p for p in read_roster(path)}
    except ValueError as e:
        raise SystemExit(str(e))


def cmd_roster(base: Path, cfg: dict, action: str, pid: str, name: str, dept: str) -> int:
    """名簿の表示・追加・削除（表計算アプリがなくても名簿を扱えるように）。"""
    path = base / cfg["roster"]
    people = list(_roster(base, cfg).values()) if path.exists() else []
    if action == "list":
        for p in people:
            print(f"{p.id}\t{p.name}\t{p.dept}")
        return 0
    pid = pid.strip()
    if action == "add":
        name = name.strip()
        if not valid_id(pid):
            print("IDは英数字・ハイフン・アンダーバーで入力してください（例：E0005）")
            return 1
        if not name:
            print("氏名を入力してください")
            return 1
        if any(p.id == pid for p in people):
            print(f"ID {pid} はすでに名簿にあります")
            return 1
        people.append(Person(pid, name, dept.strip()))
        write_roster(path, people)
        print(f"追加しました：{pid} {name}")
        return 0
    if action == "remove":
        kept = [p for p in people if p.id != pid]
        if len(kept) == len(people):
            print(f"ID {pid} は名簿にありません")
            return 1
        write_roster(path, kept)
        print(f"名簿から外しました：{pid}")
        return 0
    if action == "fixfile":   # 文字化けしない形式（BOM付きUTF-8）で保存し直す
        write_roster(path, people)
        return 0
    return 1


def _sheet_for(base: Path, cfg: dict, pid: str):
    roster = _roster(base, cfg)
    for folder, person, sheet in _answer_sheets(base / cfg["output_dir"], roster):
        if person.id == pid:
            return folder, person, sheet
    raise SystemExit(f"ID {pid} の回答が見つかりません")


def cmd_items(base: Path, cfg: dict, pid: str) -> int:
    """未確定の項目を「項目<TAB>質問<TAB>状態<TAB>読取値<TAB>選択肢1|…|4」で出力する（Macアプリ用）。"""
    from openpyxl import load_workbook
    folder, person, sheet = _sheet_for(base, cfg, pid)
    ws = load_workbook(folder / "回答確認.xlsx", data_only=True)["回答"]
    scores = {r[0]: r[4] for r in ws.iter_rows(min_row=4, values_only=True) if r[0]}
    for item in sheet.unresolved:
        opts = "|".join(SECTION_BY_KEY[item[0]].options)
        print(f"{item}\t{item_label(item)}\t{sheet.statuses.get(item, '未読取')}\t{scores.get(item) or ''}\t{opts}")
    return 0


def cmd_fix(base: Path, cfg: dict, pid: str, item: str, answer: str) -> int:
    """1項目の回答を確定する（answer は 1〜4、または none=未回答）。"""
    folder, person, _ = _sheet_for(base, cfg, pid)
    value = None if answer in ("", "none", "0") else int(answer)
    set_answer(folder, person, item, value)
    print(f"{person.id} {item} を {value if value else '未回答'} で確定しました")
    return 0


def cmd_forms(base: Path, cfg: dict, separate: bool) -> int:
    people = list(_roster(base, cfg).values())
    for path in make_forms(people, base / cfg["forms_dir"], combined=not separate):
        log.info("調査票を作成: %s", path.relative_to(base))
    log.info("%d名分。両面印刷せず片面で印刷してください（1人2枚）。", len(people))
    return 0


def _unique(dest: Path) -> Path:
    n = 2
    candidate = dest
    while candidate.exists():
        candidate = dest.with_name(f"{dest.stem}_{n}{dest.suffix}")
        n += 1
    return candidate


SCAN_SUFFIXES = {".pdf"} | IMAGE_SUFFIXES


def scan_inbox(inbox: Path) -> tuple[list[Path], list[Path]]:
    """inbox（サブフォルダも含む）から (読み取れるファイル, 対象外のファイル) を返す。

    フォルダごと入れられた場合や、スキャナが画像(JPEG/PNG/TIFF)で保存した場合にも対応する。
    権限がなく中を見られないときは PermissionError を送出する。
    """
    readable, other = [], []
    for f in sorted(inbox.rglob("*")):
        if not f.is_file() or f.name.startswith("."):    # .DS_Store や ._ で始まる隠しファイルは無視
            continue
        (readable if f.suffix.lower() in SCAN_SUFFIXES else other).append(f)
    if not readable and not other:
        list(inbox.iterdir())   # 空なのか、見る権限がないのかを区別する（権限がなければここで例外）
    return readable, other


def cmd_inbox(base: Path, cfg: dict) -> int:
    """inbox の中身を表示する。0=読み取れるファイルあり / 2=なし / 3=フォルダを見る権限がない"""
    inbox = base / cfg["input_dir"]
    inbox.mkdir(parents=True, exist_ok=True)
    print(f"場所：{inbox}")
    try:
        readable, other = scan_inbox(inbox)
    except PermissionError:
        print("このフォルダの中を見る権限がありません。")
        return 3
    print(f"読み取れるファイル（PDF・JPEG・PNG・TIFF）：{len(readable)}件")
    for f in readable[:10]:
        print(f"　・{f.relative_to(inbox)}")
    if other:
        print(f"読み取れない形式のファイル：{len(other)}件")
        for f in other[:10]:
            print(f"　・{f.relative_to(inbox)}")
        if any(f.suffix.lower() in (".heic", ".heif") for f in other):
            print("※ HEIC（iPhoneの写真）は読めません。スキャンは PDF か JPEG で保存してください。")
    return 0 if readable else 2


def read_inbox(base: Path, cfg: dict, roster: dict[str, Person]) -> int:
    inbox, output = base / cfg["input_dir"], base / cfg["output_dir"]
    inbox.mkdir(parents=True, exist_ok=True)
    errors = 0
    try:
        files, others = scan_inbox(inbox)
    except PermissionError:
        log.error("inbox の中を見る権限がありません: %s", inbox)
        return 1
    for f in others:
        log.warning("読み取れない形式のため飛ばしました: %s（PDF・JPEG・PNG・TIFF に対応）", f.relative_to(inbox))
    for pdf in files:
        try:
            pages = read_pdf(pdf, cfg.get("omr", {}))
        except Exception:
            log.exception("PDFを開けません: %s", pdf.name)
            errors += 1
            continue
        failed = False
        for pr in pages:
            if pr.error:
                dest = _unique(output / "_読取不可" / f"{pdf.stem}_p{pr.source_page}.pdf")
                save_page_pdf(pdf, pr.source_page, dest)
                log.warning("%s %dページ目: %s → %s", pdf.name, pr.source_page, pr.error, dest.relative_to(base))
                failed = True
                continue
            person = roster.get(pr.person_id) or Person(pr.person_id, "名簿なし")
            if pr.person_id not in roster:
                log.warning("ID %s が名簿にありません（%s %dページ目）", pr.person_id, pdf.name, pr.source_page)
            folder = person_folder(output, person)
            save_page_pdf(pdf, pr.source_page, folder / f"調査票スキャン_{pr.form_page}枚目.pdf")
            update_answers(folder, person, pr.items)
            todo = sum(i.status != "OK" for i in pr.items)
            log.info("読取: %s %dページ目 → %s（%d枚目, 要確認%d件）", pdf.name, pr.source_page,
                     folder.relative_to(base), pr.form_page, todo)
        if failed:
            errors += 1
        if cfg.get("move_after_processing", True):
            processed = base / cfg["processed_dir"]
            processed.mkdir(parents=True, exist_ok=True)
            shutil.move(str(pdf), _unique(processed / pdf.name))
    return errors


def _answer_sheets(output: Path, roster: dict[str, Person]):
    """output 内の本人フォルダごとに回答確認シートを読む。"""
    known = {person_folder(output, p).name: p for p in roster.values()}
    if not output.exists():
        return
    for folder in sorted(p for p in output.iterdir() if p.is_dir() and not p.name.startswith("_")):
        person = known.get(folder.name)
        if person is None:
            pid, _, name = folder.name.partition("_")
            person = Person(pid, name)
        sheet = load_answer_sheet(folder, person)
        if sheet is not None:
            yield folder, person, sheet


def list_todo(base: Path, cfg: dict, roster: dict[str, Person]) -> tuple[list[dict], list[Path]]:
    """要確認の人（回答確認シートに未確定の項目がある人）と、読取不可のページを返す。"""
    output = base / cfg["output_dir"]
    people = []
    for folder, person, sheet in _answer_sheets(output, roster):
        if sheet.unresolved:
            people.append({
                "person": person, "folder": folder,
                "items": [(i, sheet.statuses.get(i, "未読取")) for i in sheet.unresolved],
                "scans": sorted(folder.glob("調査票スキャン_*.pdf")),
            })
    unreadable = sorted((output / "_読取不可").glob("*.pdf")) if (output / "_読取不可").exists() else []
    return people, unreadable


def cmd_todo(base: Path, cfg: dict, tsv: bool) -> int:
    people, unreadable = list_todo(base, cfg, _roster(base, cfg))
    if tsv:   # Macアプリ用: 種別<TAB>開くファイル(|区切り)<TAB>表示名
        for p in people:
            counts: dict[str, int] = {}
            for _, st in p["items"]:
                counts[st] = counts.get(st, 0) + 1
            detail = "・".join(f"{k}{v}件" for k, v in counts.items())
            files = [p["folder"] / "回答確認.xlsx", *p["scans"]]
            print(f"PERSON\t{'|'.join(map(str, files))}\t{p['person'].id} {p['person'].name}（{detail}）")
        for f in unreadable:
            print(f"UNREADABLE\t{f}\t読取不可のページ：{f.name}")
        return 0
    if not people and not unreadable:
        print("要確認の人はいません。")
        return 0
    for p in people:
        items = "、".join(f"{i}（{st}）" for i, st in p["items"])
        print(f"■ {p['person'].id} {p['person'].name}：{len(p['items'])}件 … {items}")
        print(f"    修正するファイル: {p['folder'] / '回答確認.xlsx'}")
    for f in unreadable:
        print(f"■ 読取不可のページ: {f}")
    return 0


def build_outputs(base: Path, cfg: dict, roster: dict[str, Person]) -> None:
    output = base / cfg["output_dir"]
    results = []
    for folder, person, sheet in _answer_sheets(output, roster):
        for old in folder.glob("結果通知書_*.pdf"):
            old.unlink()
        if sheet.unresolved:
            log.warning("%s: 要確認 %d件のため結果通知書は未作成（%s を修正してください）",
                        person.name, len(sheet.unresolved), (folder / "回答確認.xlsx").relative_to(base))
            results.append((sheet, None))
            continue
        res = score(sheet.answers, cfg["criteria"])
        write_report(folder / f"結果通知書_{safe_name(person.id)}_{safe_name(person.name)}.pdf", person, res, cfg)
        results.append((sheet, res))

    summary = output / cfg["summary_excel"]
    write_summary(summary, results, cfg)
    reports = sorted(output.glob("*/結果通知書_*.pdf"))
    bundle = output / "結果通知書_印刷用_全員分.pdf"
    bundle.unlink(missing_ok=True)
    if reports:
        combine(reports, bundle)
    done = sum(r is not None for _, r in results)
    log.info("集計: %s（%d名中 %d名の結果確定）", summary.relative_to(base), len(results), done)
    if reports:
        log.info("印刷用: %s（%d枚）", bundle.relative_to(base), len(reports))
    missing = sorted(set(roster) - {s.person.id for s, _ in results})
    if missing:
        log.info("未提出: %d名（%s）", len(missing), ", ".join(missing[:20]) + (" …" if len(missing) > 20 else ""))


def cmd_run(base: Path, cfg: dict) -> int:
    roster = _roster(base, cfg)
    (base / cfg["output_dir"]).mkdir(parents=True, exist_ok=True)
    errors = read_inbox(base, cfg, roster)
    build_outputs(base, cfg, roster)
    return 1 if errors else 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="pdfeval", description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("command", nargs="?", default="run", choices=["run", "forms", "todo", "inbox", "items", "fix",
                                 "roster-list", "roster-add", "roster-remove", "roster-fixfile"])
    parser.add_argument("--config", type=Path, default=BASE_DIR / "config.json")
    parser.add_argument("--base", type=Path, default=None,
                        help="名簿・inbox・output を置く基準フォルダ (既定: 設定ファイルと同じ場所)")
    parser.add_argument("--separate", action="store_true", help="forms: 1人1ファイルで出力する")
    parser.add_argument("--tsv", action="store_true", help="todo: 機械向けのタブ区切りで出力する")
    parser.add_argument("--id", default="", help="items/fix/roster-add/roster-remove: 対象のID")
    parser.add_argument("--item", default="", help="fix: 項目（例 A5）")
    parser.add_argument("--answer", default="", help="fix: 回答 1〜4、または none")
    parser.add_argument("--name", default="", help="roster-add: 氏名")
    parser.add_argument("--dept", default="", help="roster-add: 所属")
    parser.add_argument("--version", action="version", version=__version__)
    args = parser.parse_args(argv)

    base = (args.base or args.config.resolve().parent).resolve()
    (base / "logs").mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s",
        handlers=[logging.StreamHandler(),
                  logging.FileHandler(base / "logs" / "pdfeval.log", encoding="utf-8")],
        force=True,
    )
    cfg = load_config(args.config)
    if args.command == "forms":
        return cmd_forms(base, cfg, args.separate)
    if args.command == "todo":
        return cmd_todo(base, cfg, args.tsv)
    if args.command == "inbox":
        return cmd_inbox(base, cfg)
    if args.command == "items":
        return cmd_items(base, cfg, args.id)
    if args.command == "fix":
        return cmd_fix(base, cfg, args.id, args.item, args.answer)
    if args.command.startswith("roster-"):
        return cmd_roster(base, cfg, args.command[len("roster-"):], args.id, args.name, args.dept)
    return cmd_run(base, cfg)
