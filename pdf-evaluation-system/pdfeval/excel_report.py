"""全員分の結果を集計Excelに書き出す（円グラフ付き）。実施者のみが扱うデータです。"""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .organizer import AnswerSheet
from .questionnaire import ALL_ITEMS, DOMAINS, SCALES, Result, item_label, scale_level

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
HIGH_FILL = PatternFill("solid", fgColor="F8CBAD")
TODO_FILL = PatternFill("solid", fgColor="FFE699")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _table(ws: Worksheet, row: int, col: int, headers: list, rows: list[list]) -> None:
    for j, h in enumerate(headers):
        c = ws.cell(row=row, column=col + j, value=h)
        c.font, c.fill, c.border = Font(bold=True), HEADER_FILL, BORDER
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, values in enumerate(rows, start=1):
        for j, v in enumerate(values):
            c = ws.cell(row=row + i, column=col + j, value=v)
            c.border = BORDER


def _pie(ws: Worksheet, title: str, cat_col: int, val_col: int, first: int, last: int,
         anchor: str, colors: list[str]) -> None:
    pie = PieChart()
    pie.title = title
    pie.add_data(Reference(ws, min_col=val_col, min_row=first - 1, max_row=last), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=cat_col, min_row=first, max_row=last))
    for i, color in enumerate(colors):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = color
        pie.series[0].dPt.append(pt)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showCatName = True
    pie.dataLabels.showVal = pie.dataLabels.showSerName = pie.dataLabels.showLegendKey = False
    pie.height, pie.width = 7.5, 10
    ws.add_chart(pie, anchor)


def write_summary(path: Path, results: list[tuple[AnswerSheet, Result | None]], cfg: dict) -> None:
    th = cfg.get("scale_thresholds", {})
    wb = Workbook()

    # --- 集計 ---
    ws = wb.active
    ws.title = "集計"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws["A1"] = "ストレスチェック 集計"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "※個人の結果を含みます。実施者・実施事務従事者以外に開示しないでください。"
    ws["A2"].font = Font(color="C00000")

    done = [(s, r) for s, r in results if r is not None]
    verdict = Counter("該当" if r.high_stress else "非該当" for _, r in done)
    verdict_rows = [["高ストレス者に該当", verdict["該当"]], ["該当しない", verdict["非該当"]],
                    ["結果未確定（要確認あり）", len(results) - len(done)]]
    _table(ws, 4, 1, ["判定", "人数"], verdict_rows)
    ws.cell(row=8, column=1, value="受検者数").font = Font(bold=True)
    ws.cell(row=8, column=2, value=len(results))
    ws.cell(row=9, column=1, value="高ストレス者の割合").font = Font(bold=True)
    ws.cell(row=9, column=2, value=(verdict["該当"] / len(done)) if done else 0).number_format = "0.0%"
    _pie(ws, "高ストレス判定", 1, 2, 5, 7, "D4", ["C8553D", "2E8B7A", "BFBFBF"])

    levels = Counter(scale_level(r.b / DOMAINS["B"][1], th) for _, r in done)
    lv_rows = [[k, levels[k]] for k in ("良好", "ふつう", "要注意")]
    _table(ws, 20, 1, ["心身のストレス反応", "人数"], lv_rows)
    _pie(ws, "心身のストレス反応の分布", 1, 2, 21, 23, "D20", ["2E8B7A", "C9A227", "C8553D"])

    # 所属別（集団分析の目安。人数が少ない所属は個人が特定されるおそれがあるので注意）
    by_dept: dict[str, list[Result]] = {}
    for s, r in done:
        by_dept.setdefault(s.person.dept or "（所属なし）", []).append(r)
    dept_rows = [[d, len(rs), sum(r.high_stress for r in rs),
                  round(sum(r.a for r in rs) / len(rs), 1), round(sum(r.b for r in rs) / len(rs), 1),
                  round(sum(r.c for r in rs) / len(rs), 1)] for d, rs in sorted(by_dept.items())]
    _table(ws, 36, 1, ["所属", "人数", "高ストレス者", "平均A", "平均B", "平均C"], dept_rows)
    ws.cell(row=35, column=1, value="所属別（10人未満の集団は個人が特定されやすいため取扱注意）").font = Font(bold=True)
    for col, w in {"A": 26, "B": 10, "C": 12, "D": 9, "E": 9, "F": 9}.items():
        ws.column_dimensions[col].width = w

    # --- 一覧 ---
    wl = wb.create_sheet("一覧")
    headers = ["ID", "氏名", "所属", "状態", "A 仕事のストレス要因", "B 心身のストレス反応",
               "C 周囲のサポート", "A+C", "判定", "判定理由"] + [name for _, name, _ in SCALES]
    rows = []
    for s, r in results:
        base = [s.person.id, s.person.name, s.person.dept,
                "確定" if r else f"要確認 {len(s.unresolved)}件"]
        if r:
            rows.append(base + [r.a, r.b, r.c, r.a + r.c, "該当" if r.high_stress else "非該当", r.reason]
                        + [avg for _, _, avg in r.scales])
        else:
            rows.append(base + [None] * (len(headers) - 4))
    _table(wl, 1, 1, headers, rows)
    for i, (s, r) in enumerate(results, start=2):
        fill = TODO_FILL if r is None else (HIGH_FILL if r.high_stress else None)
        if fill:
            for c in range(1, 10):
                wl.cell(row=i, column=c).fill = fill
    wl.row_dimensions[1].height = 45
    wl.freeze_panes = "D2"
    wl.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for c in range(1, len(headers) + 1):
        wl.column_dimensions[get_column_letter(c)].width = 10
    wl.column_dimensions["B"].width = 14
    wl.column_dimensions["J"].width = 40

    # --- 回答データ ---
    wd = wb.create_sheet("回答データ")
    _table(wd, 1, 1, ["ID", "氏名"] + ALL_ITEMS,
           [[s.person.id, s.person.name] + [s.answers.get(i) for i in ALL_ITEMS] for s, _ in results])
    wd.freeze_panes = "C2"

    # --- 要確認 ---
    wt = wb.create_sheet("要確認")
    todo = [[s.person.id, s.person.name, i, item_label(i), s.statuses.get(i, "未読取"),
             str(s.folder / "回答確認.xlsx")] for s, _ in results for i in s.unresolved]
    _table(wt, 1, 1, ["ID", "氏名", "項目", "質問", "状態", "修正するファイル"], todo)
    for col, w in {"A": 8, "B": 14, "C": 6, "D": 40, "E": 9, "F": 60}.items():
        wt.column_dimensions[col].width = w

    wb.save(path)
