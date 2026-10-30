"""python -m unittest discover tests"""
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "sample"))

from pdfeval.main import build_outputs, load_config, read_inbox, _roster  # noqa: E402
from pdfeval.questionnaire import ALL_ITEMS, score  # noqa: E402

CONFIG = load_config(ROOT / "config.json")


class ScoringTest(unittest.TestCase):
    def test_item_count(self):
        self.assertEqual(len(ALL_ITEMS), 57)

    def test_lowest_and_highest_stress(self):
        # 逆転項目を考慮して「最もストレスが低い」「最も高い」回答を作る
        from pdfeval.questionnaire import REVERSED
        low = {i: (4 if i in REVERSED else 1) for i in ALL_ITEMS}
        high = {i: (1 if i in REVERSED else 4) for i in ALL_ITEMS}
        r = score(low, CONFIG["criteria"])
        self.assertEqual((r.a, r.b, r.c, r.high_stress), (17, 29, 9, False))
        r = score(high, CONFIG["criteria"])
        self.assertEqual((r.a, r.b, r.c, r.high_stress), (68, 116, 36, True))

    def test_criteria_boundaries(self):
        crit = CONFIG["criteria"]
        def answers_with(b_total, ac_high):
            from pdfeval.questionnaire import REVERSED
            ans = {i: (1 if i in REVERSED else 4) if ac_high else (4 if i in REVERSED else 1)
                   for i in ALL_ITEMS if i[0] in "ACD"}
            pts = [1] * 29
            for k in range(b_total - 29):
                pts[k % 29] += 1
            for n, p in enumerate(pts, start=1):
                ans[f"B{n}"] = 5 - p if n <= 3 else p
            return ans
        self.assertFalse(score(answers_with(76, False), crit).high_stress)
        self.assertTrue(score(answers_with(77, False), crit).high_stress)   # 条件①
        self.assertFalse(score(answers_with(62, True), crit).high_stress)
        self.assertTrue(score(answers_with(63, True), crit).high_stress)    # 条件②


class EndToEndTest(unittest.TestCase):
    def test_scan_to_reports(self):
        import simulate_scan

        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            shutil.copy(ROOT / "sample" / "名簿.csv", base / "名簿.csv")
            orig_root = simulate_scan.ROOT
            simulate_scan.ROOT = base
            (base / "sample").mkdir()
            shutil.copy(ROOT / "sample" / "名簿.csv", base / "sample" / "名簿.csv")
            try:
                simulate_scan.main()
            finally:
                simulate_scan.ROOT = orig_root
            truth = json.loads((base / "sample" / "正解.json").read_text(encoding="utf-8"))

            roster = _roster(base, CONFIG)
            self.assertEqual(read_inbox(base, CONFIG, roster), 0)
            build_outputs(base, CONFIG, roster)

            out = base / "output"
            for pid, answers in truth.items():
                folder = next(out.glob(f"{pid}_*"))
                ws = load_workbook(folder / "回答確認.xlsx")["回答"]
                read = {r[0]: (r[2], r[3]) for r in ws.iter_rows(min_row=4, values_only=True)}
                for item, expected in answers.items():
                    if expected is None:
                        self.assertIn(read[item][1], ("未記入", "複数"), f"{pid} {item}")
                    else:
                        self.assertEqual(read[item], (expected, "OK"), f"{pid} {item}")
            # 要確認が残る E0003 だけ結果通知書がない
            self.assertEqual(len(list(out.glob("*/結果通知書_*.pdf"))), 3)
            self.assertFalse(list(out.glob("E0003_*/結果通知書_*.pdf")))
            wb = load_workbook(out / "ストレスチェック集計.xlsx")
            self.assertEqual(wb.sheetnames, ["集計", "一覧", "回答データ", "要確認"])
            self.assertEqual(wb["要確認"].max_row, 3)
            self.assertEqual(len(wb["集計"]._charts), 2)


if __name__ == "__main__":
    unittest.main()


class SettingsTest(unittest.TestCase):
    def test_list_and_set(self):
        import json
        from pdfeval.settings import list_settings, set_setting
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "config.json"
            shutil.copy(ROOT / "config.json", path)
            cfg = json.loads(path.read_text(encoding="utf-8"))
            current = {k: v for k, _, v, _ in list_settings(cfg)}
            self.assertEqual(current["contact"], "株式会社proceed YADOLI訪問看護ステーション")   # 頭の「相談先：」は除く
            self.assertEqual(current["contact_detail"], "")
            # 頭の文字を自分で入れても二重にならない／改行は取り除く
            self.assertEqual(set_setting(path, "contact", "相談先：YADOLI 相談窓口\n"), "相談先：YADOLI 相談窓口")
            self.assertEqual(set_setting(path, "contact_detail", "電話 03-1234-5678"), "電話 03-1234-5678")
            saved = json.loads(path.read_text(encoding="utf-8"))
            self.assertEqual(saved["criteria"], cfg["criteria"])        # ほかの設定は変わらない
            with self.assertRaises(ValueError):
                set_setting(path, "criteria", "x")

    def test_old_placeholder_is_migrated(self):
        import json
        from pdfeval.main import load_config
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "config.json"
            cfg = json.loads((ROOT / "config.json").read_text(encoding="utf-8"))
            cfg["report"]["contact"] = "相談窓口：〇〇部 健康管理担当（内線 0000）"
            del cfg["report"]["contact_detail"]
            path.write_text(json.dumps(cfg, ensure_ascii=False), encoding="utf-8")
            new = load_config(path)
            self.assertEqual(new["report"]["contact"], "相談先：株式会社proceed YADOLI訪問看護ステーション")
            self.assertEqual(new["report"]["contact_detail"], "")


class GroupReportTest(unittest.TestCase):
    def _results(self, n, dept="訪問看護"):
        import random

        from pdfeval.organizer import AnswerSheet
        from pdfeval.forms import Person
        from pdfeval.questionnaire import ALL_ITEMS, score
        rng = random.Random(n)
        out = []
        for i in range(n):
            ans = {k: rng.randint(1, 4) for k in ALL_ITEMS}
            out.append((AnswerSheet(Person(f"E{i}", "x", dept), Path("."), ans, {}), score(ans, CONFIG["criteria"])))
        return out

    def test_one_page_for_various_sizes(self):
        import pypdfium2 as pdfium

        from pdfeval.group_report import write_group_report
        from pdfeval.organizer import AnswerSheet
        from pdfeval.forms import Person
        with tempfile.TemporaryDirectory() as tmp:
            cases = {
                "zero": [(AnswerSheet(Person("E1", "x", "y"), Path("."), {}, {}), None)],
                "small": self._results(4),
                "big": self._results(15) + self._results(6, "事務") + self._results(3, "リハビリ"),
            }
            for name, results in cases.items():
                path = write_group_report(Path(tmp) / f"{name}.pdf", results, 30, CONFIG)
                doc = pdfium.PdfDocument(str(path))
                self.assertEqual(len(doc), 1, name)         # 必ずA4・1枚
                text = doc[0].get_textpage().get_text_range()
                doc.close()
                if name == "small":
                    self.assertIn("10名未満", text)
                if name == "big":
                    self.assertIn("人数が少ないため非表示", text)   # 6名・3名の所属は数値を出さない

    def test_judge_scale_scoring(self):
        from pdfeval.group_report import _judge_score
        # 「そうだ(1)」ばかり → 量的負担は最大の12点、「ちがう(4)」ばかり → 最小の3点
        self.assertEqual(_judge_score({"A1": 1, "A2": 1, "A3": 1}, ("A1", "A2", "A3")), 12)
        self.assertEqual(_judge_score({"A1": 4, "A2": 4, "A3": 4}, ("A1", "A2", "A3")), 3)
        self.assertIsNone(_judge_score({"A1": 4, "A2": None, "A3": 4}, ("A1", "A2", "A3")))


class JudgmentChartTest(unittest.TestCase):
    def test_health_risk(self):
        import math

        from pdfeval.group_report import health_risk
        self.assertIsNone(health_risk(None, 7, 7))
        self.assertIsNone(health_risk({"mean_x": 8, "mean_y": None, "coef_x": 0.1, "coef_y": -0.1}, 7, 7))
        p = {"mean_x": 8.0, "mean_y": 7.0, "coef_x": 0.1, "coef_y": -0.1}
        self.assertAlmostEqual(health_risk(p, 8.0, 7.0), 100.0)             # 全国平均と同じなら100
        self.assertAlmostEqual(health_risk(p, 9.0, 7.0), 100 * math.exp(0.1))

    def test_report_with_and_without_values(self):
        import copy

        import pypdfium2 as pdfium

        from pdfeval.group_report import write_group_report
        results = GroupReportTest()._results(15)
        cfg = copy.deepcopy(CONFIG)
        with tempfile.TemporaryDirectory() as tmp:
            plain = write_group_report(Path(tmp) / "a.pdf", results, 15, cfg)
            cfg["judgment_chart"] = {"quantity_control": {"mean_x": 8, "mean_y": 7.5, "coef_x": 0.08, "coef_y": -0.09},
                                     "support": {"mean_x": 7.5, "mean_y": 8, "coef_x": -0.1, "coef_y": -0.1}}
            filled = write_group_report(Path(tmp) / "b.pdf", results, 15, cfg)
            for path, expect in ((plain, "設定に厚生労働省"), (filled, "総合健康リスク")):
                doc = pdfium.PdfDocument(str(path))
                self.assertEqual(len(doc), 1)
                self.assertIn(expect, doc[0].get_textpage().get_text_range())
                doc.close()

    def test_numeric_setting(self):
        import json

        from pdfeval.settings import set_setting
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "config.json"
            shutil.copy(ROOT / "config.json", path)
            self.assertEqual(set_setting(path, "judgment_chart/support/coef_x", "－０．１"), "-0.1")
            self.assertEqual(json.loads(path.read_text())["judgment_chart"]["support"]["coef_x"], -0.1)
            with self.assertRaises(ValueError):
                set_setting(path, "judgment_chart/support/coef_x", "abc")


class CompanyGroupingTest(unittest.TestCase):
    def test_group_mapping_and_label(self):
        import copy
        import json

        import pypdfium2 as pdfium

        from pdfeval.group_report import write_group_report
        from pdfeval.settings import group_of, set_group, set_setting
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "config.json"
            shutil.copy(ROOT / "config.json", path)
            set_group(path, "リハビリ", "医療部")
            set_group(path, "訪問看護", " 医療部 ")
            set_setting(path, "company_name", "株式会社テスト")
            set_setting(path, "group_label", "部署")
            cfg = json.loads(path.read_text(encoding="utf-8"))
            self.assertEqual(group_of("リハビリ", cfg), "医療部")
            self.assertEqual(group_of("事務", cfg), "事務")
            self.assertEqual(group_of("", cfg), "（部署なし）")
            set_group(path, "リハビリ", "")                       # 空欄で元に戻る
            self.assertNotIn("リハビリ", json.loads(path.read_text(encoding="utf-8"))["group_map"])

            # 8名＋6名は単独では非表示だが、まとめると14名で数値が出る
            cfg = copy.deepcopy(CONFIG)
            cfg["report"].update(company_name="株式会社テスト", group_label="部署")
            cfg["group_map"] = {"リハビリ": "医療部", "訪問看護": "医療部"}
            t = GroupReportTest()
            results = t._results(8, "リハビリ") + t._results(6, "訪問看護")
            pdf = write_group_report(Path(tmp) / "g.pdf", results, 14, cfg)
            doc = pdfium.PdfDocument(str(pdf))
            text = doc[0].get_textpage().get_text_range()
            doc.close()
            self.assertIn("株式会社テスト", text)
            self.assertIn("部署別", text)
            self.assertIn("医療部", text)
            self.assertNotIn("人数が少ないため非表示", text)
